from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass(slots=True, frozen=True)
class ExcelAnswerRow:
    question_number: int
    question: str
    answer: str
    status: str
    confidence: float | None = None
    evidence_ids: tuple[str, ...] = ()
    pages: tuple[int, ...] = ()
    error: str | None = None


class ExcelExporter:
    def __init__(
        self,
        output_directory: str | Path = "outputs",
    ) -> None:
        self.output_directory = Path(output_directory)
        self.output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

    def export_answers(
        self,
        *,
        company_name: str,
        report_year: int,
        rows: Sequence[ExcelAnswerRow],
        report_id: str | None = None,
        report_url: str | None = None,
    ) -> Path:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Cevaplar"

        worksheet["A1"] = "Şirket"
        worksheet["B1"] = company_name
        worksheet["A2"] = "Rapor Yılı"
        worksheet["B2"] = report_year
        worksheet["A3"] = "Rapor Kimliği"
        worksheet["B3"] = report_id or ""
        worksheet["A4"] = "Rapor URL"
        worksheet["B4"] = report_url or ""

        for cell in ("A1", "A2", "A3", "A4"):
            worksheet[cell].font = Font(bold=True)

        header_row = 6

        headers = (
            "Soru No",
            "Soru",
            "Cevap",
            "Durum",
            "Güven",
            "Kanıt Kimlikleri",
            "Sayfalar",
            "Hata",
        )

        for column_index, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=header_row,
                column=column_index,
                value=header,
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row_index, answer_row in enumerate(
            rows,
            start=header_row + 1,
        ):
            values = (
                answer_row.question_number,
                answer_row.question,
                answer_row.answer,
                answer_row.status,
                answer_row.confidence,
                ", ".join(answer_row.evidence_ids),
                ", ".join(
                    str(page)
                    for page in answer_row.pages
                ),
                answer_row.error or "",
            )

            for column_index, value in enumerate(
                values,
                start=1,
            ):
                cell = worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        column_widths = {
            "A": 10,
            "B": 55,
            "C": 70,
            "D": 22,
            "E": 12,
            "F": 35,
            "G": 18,
            "H": 45,
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        worksheet.freeze_panes = f"A{header_row + 1}"
        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"{get_column_letter(len(headers))}"
            f"{header_row + len(rows)}"
        )

        safe_company_name = self._safe_filename(
            company_name,
        )

        output_path = self.output_directory / (
            f"{safe_company_name}_{report_year}_answers.xlsx"
        )

        workbook.save(output_path)

        return output_path

    @staticmethod
    def _safe_filename(value: str) -> str:
        translation_table = str.maketrans(
            {
                "ı": "i",
                "İ": "I",
                "ş": "s",
                "Ş": "S",
                "ğ": "g",
                "Ğ": "G",
                "ü": "u",
                "Ü": "U",
                "ö": "o",
                "Ö": "O",
                "ç": "c",
                "Ç": "C",
            },
        )

        normalized = value.translate(
            translation_table,
        )

        normalized = re.sub(
            r"[^a-zA-Z0-9_-]+",
            "_",
            normalized,
        )

        return normalized.strip("_").lower()


__all__ = [
    "ExcelAnswerRow",
    "ExcelExporter",
]